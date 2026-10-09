"""文档解析：PDF / DOCX / PPTX / TXT / Markdown / 图片(OCR可选)。

解析产出全文 + 章节结构 outline（用于复习知识树与引用定位），再按章节分块。
"""
import logging
import os
import re
from dataclasses import dataclass, field

logger = logging.getLogger("zhinan.parser")

SUPPORTED_EXT = {".pdf", ".doc", ".docx", ".ppt", ".pptx",
                 ".txt", ".md", ".markdown", ".text",
                 ".jpg", ".jpeg", ".png"}


@dataclass
class ParsedDoc:
    text: str = ""
    outline: list[dict] = field(default_factory=list)   # [{title, page, section}]
    pages: dict[int, str] = field(default_factory=dict)  # page -> text
    ocr_used: bool = False

    def page_of(self, offset: int) -> int:
        """根据全文偏移定位页码。"""
        acc = 0
        for page in sorted(self.pages):
            acc += len(self.pages[page])
            if offset < acc:
                return page
        return max(self.pages) if self.pages else 1


def parse_pdf(path: str) -> ParsedDoc:
    import pymupdf  # PyMuPDF 新包名

    doc = ParsedDoc()
    total_chars = 0
    try:
        with pymupdf.open(path) as pdf:
            total_pages = len(pdf)
            for i, page in enumerate(pdf, start=1):
                txt = page.get_text("text").strip()
                if not txt and _ocr_available():
                    txt = _run_ocr_page(page)
                    doc.ocr_used = True
                doc.pages[i] = txt
                total_chars += len(txt)
                # 用字号识别标题行
                for block in page.get_text("dict")["blocks"]:
                    for line in block.get("lines", []):
                        sizes = [s["size"] for s in line.get("spans", []) if s["text"].strip()]
                        if not sizes:
                            continue
                        line_text = "".join(s["text"] for s in line["spans"]).strip()
                        if _looks_like_heading(line_text) and max(sizes) >= 14 and len(line_text) <= 40:
                            doc.outline.append({"title": line_text, "page": i, "section": line_text})
                if doc.ocr_used:
                    break  # OCR 模式下剩余页按图片整页处理
            if doc.ocr_used:
                for i in range(1, total_pages + 1):
                    if i not in doc.pages:
                        page = pdf[i - 1]
                        pix = page.get_pixmap(dpi=150)
                        import tempfile
                        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
                            pix.save(f.name)
                            doc.pages[i] = _run_ocr_image(f.name)
                            os.unlink(f.name)
    except Exception as exc:  # noqa: BLE001
        logger.warning("PDF 打开失败，降级为「元数据占位条目」：%s → %s", os.path.basename(path), exc)
        return metadata_fallback_parsedoc(path, "pdf")

    doc.text = "\n".join(doc.pages.get(i, "") for i in sorted(doc.pages))

    # 检测「几乎空白的扫描版 PDF」：非空白字符极少且 OCR 不可用
    non_ws = len(re.sub(r"\s+", "", doc.text))
    if non_ws < 100 and not doc.ocr_used and not _ocr_available():
        logger.warning("PDF %s 抽取文本极少（仅 %d 非空白字符），疑似扫描版且未装 OCR，降级为「元数据占位条目」。",
                       os.path.basename(path), non_ws)
        return metadata_fallback_parsedoc(path, "pdf_scan")
    return doc


def parse_docx(path: str) -> ParsedDoc:
    doc = ParsedDoc()
    try:
        from docx import Document
        d = Document(path)
    except Exception as exc:  # noqa: BLE001
        # 结构损坏的 docx：如 _rels 引用了不存在的部件（WPS 自定义 UI 残留），降级占位
        logger.warning("DOCX 打开失败（%s），降级为「元数据占位条目」：%s → %s",
                       type(exc).__name__, os.path.basename(path), exc)
        return metadata_fallback_parsedoc(path, "docx")
    page = 1
    for para in d.paragraphs:
        style = (para.style.name or "").lower()
        text = para.text.strip()
        if not text:
            continue
        if "heading" in style or "标题" in style:
            doc.outline.append({"title": text, "page": page, "section": text})
        doc.text += text + "\n"
    doc.pages[1] = doc.text
    return doc


def parse_pptx(path: str) -> ParsedDoc:
    doc = ParsedDoc()
    try:
        from pptx import Presentation
        prs = Presentation(path)
    except Exception as exc:  # noqa: BLE001
        logger.warning("PPTX 打开失败（%s），降级为「元数据占位条目」：%s → %s",
                       type(exc).__name__, os.path.basename(path), exc)
        return metadata_fallback_parsedoc(path, "pptx")
    for i, slide in enumerate(prs.slides, start=1):
        lines = []
        title = ""
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                t = "".join(run.text for run in para.runs).strip()
                if t:
                    lines.append(t)
        if lines:
            title = lines[0]
            doc.outline.append({"title": title, "page": i, "section": title})
        doc.pages[i] = "\n".join(lines)
    doc.text = "\n".join(doc.pages.get(i, "") for i in sorted(doc.pages))
    return doc


def parse_plain(path: str) -> ParsedDoc:
    doc = ParsedDoc()
    with open(path, encoding="utf-8", errors="ignore") as f:
        doc.text = f.read()
    for m in re.finditer(r"^#{1,4}\s+(.+)$", doc.text, re.M):
        doc.outline.append({"title": m.group(1).strip(), "page": 1, "section": m.group(1).strip()})
    doc.pages[1] = doc.text
    return doc


def parse_image(path: str) -> ParsedDoc:
    doc = ParsedDoc()
    if _ocr_available():
        doc.pages[1] = _run_ocr_image(path)
        doc.ocr_used = True
    else:
        doc.pages[1] = ""
    doc.text = doc.pages[1]
    if not doc.text.strip():
        # OCR 无文字或 OCR 未安装 → 方案 C：以文件名 + 路径元数据作为一段可检索文本
        logger.warning("图片 %s 未提取到文字，降级为「元数据占位条目」入库。",
                       os.path.basename(path))
        return metadata_fallback_parsedoc(path, "image")
    return doc


# -------------------- 元数据占位条目（方案 C 的降级兜底） --------------------
def metadata_fallback_parsedoc(path: str, kind: str) -> ParsedDoc:
    """为无法完整解析的文件生成一段元数据描述，保证至少能被关键词命中。

    生成的文本包含：课程名、子分类（期末试卷/课件/…）、文件名、大小、格式说明。
    outline 第一项用文件名做标题，分块器会按「元数据 + 章节」切出 1~2 个可检索 chunk。
    """
    try:
        size_bytes = os.path.getsize(path)
        size_mb = size_bytes / 1024 / 1024
        size_str = f"{size_mb:.2f}MB" if size_mb >= 1 else f"{size_bytes/1024:.1f}KB"
    except OSError:
        size_bytes, size_str = 0, "未知"
    folder_parts = [p for p in os.path.normpath(path).split(os.sep) if p]
    try:
        idx = folder_parts.index("学习资料")
        course = folder_parts[idx + 1] if idx + 1 < len(folder_parts) else ""
        subcat = folder_parts[idx + 2] if idx + 2 < len(folder_parts) else ""
    except ValueError:
        course, subcat = "", ""
    filename = os.path.basename(path)
    ext = os.path.splitext(filename)[1].lstrip(".").upper()

    kind_label = {
        "doc": "Microsoft Word 97-2003 文档(.doc)",
        "ppt": "Microsoft PowerPoint 97-2003 演示文档(.ppt)",
        "image": "图片资料（JPG/PNG 扫描件或讲义照片）",
        "pdf": "PDF 文档（打开失败或结构损坏）",
        "pdf_scan": "扫描版 PDF（纯图片页，当前环境缺少 OCR 引擎无法提取文字）",
        "docx": "Word 文档(.docx，文件结构损坏或自定义扩展部件缺失，无法正常读取正文)",
        "pptx": "PowerPoint 文档(.pptx，文件结构损坏或自定义扩展部件缺失，无法正常读取正文)",
    }.get(kind, f"{kind} 文档")

    course_line = f"所属课程：{course}" if course else "所属课程：未分类"
    sub_line = f"资料分类：{subcat}" if subcat else "资料分类：学习资料"

    # 推断 doc_type
    name_lower = filename.lower()
    doc_type = "学习资料"
    if any(k in filename for k in ("期末", "试卷", "A卷", "B卷", "VC201", "print")):
        doc_type = "期末试卷 / 期末练习题"
    elif any(k in filename for k in ("答案", "解答", "解析", "参考答案")):
        doc_type = "习题参考答案 / 解析"
    elif any(k in filename for k in ("复习", "笔记", "汇总", "知识点", "总结", "提纲")):
        doc_type = "复习笔记 / 知识点汇总"
    elif any(k in name_lower for k in ("ppt", "课件", "chapter", "ch.")) or "第" in filename and "章" in filename:
        doc_type = "课程课件"
    elif any(k in filename for k in ("实验", "Experiment", "Lab")):
        doc_type = "实验报告 / 实验指导"
    elif any(k in filename for k in ("作业", "习题集")):
        doc_type = "课程作业 / 习题"

    note = (
        "说明：本文件在当前运行环境下未能抽取正文（可能是老版 .doc/.ppt 格式，或图片资料未装 OCR 引擎）。"
        "以下仅为文件元数据，作为知识库的占位条目，便于你在聊天中按文件名、课程、资料类型检索到它；"
        "若需正文内容，请把 .doc/.ppt 另存为 .docx/.pptx，或安装 pywin32 + Microsoft Office、paddleocr。"
    )

    text_lines = [
        f"【{kind_label} 占位条目】 {filename}",
        course_line,
        f"资料类型：{doc_type}",
        sub_line,
        f"原始文件名：{filename}",
        f"文件格式：{ext}",
        f"文件大小：{size_str}（{size_bytes} bytes）",
        f"课程下所属子目录：{subcat or filename}",
        note,
    ]
    text = "\n".join(text_lines)

    title = f"{filename}（{course or '未分类'}-{doc_type}）"
    return ParsedDoc(
        text=text,
        outline=[{"title": title, "level": 1, "page": 1}],
        pages={1: text},
        ocr_used=False,
    )


def _ocr_available() -> bool:
    try:
        import paddleocr  # noqa: F401
        return True
    except ImportError:
        return False


def _run_ocr_page(page) -> str:
    import tempfile
    pix = page.get_pixmap(dpi=150)
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
        pix.save(f.name)
        text = _run_ocr_image(f.name)
        os.unlink(f.name)
    return text


def _run_ocr_image(path: str) -> str:
    from paddleocr import PaddleOCR
    ocr = PaddleOCR(use_angle_cls=True, lang="ch", show_log=False)
    result = ocr.ocr(path, cls=True)
    lines = []
    for item in (result[0] or []):
        try:
            lines.append(item[1][0])
        except (TypeError, IndexError):
            continue
    return "\n".join(lines)


def _looks_like_heading(text: str) -> bool:
    return bool(re.match(r"^(第?[一二三四五六七八九十百\d]+[章节讲单元部分篇]|\d+(\.\d+)*\s+\S+|[一二三四五六七八九十]+[、.．]\S+)", text))


# -------------------- 老版 .doc / .ppt 解析（Windows + Office 优先，否则降级） --------------------
def _office_to_temp(path: str, target_ext: str) -> str:
    """用本机 Microsoft Office（通过 pywin32/win32com）把旧版 .doc/.ppt 另存为 docx/pptx。
    成功返回临时文件路径；失败抛出 ValueError。"""
    import tempfile
    ext = os.path.splitext(path)[1].lower()
    if target_ext == ".docx" and ext != ".doc":
        raise ValueError(f"仅支持 .doc 转 .docx：{ext}")
    if target_ext == ".pptx" and ext != ".ppt":
        raise ValueError(f"仅支持 .ppt 转 .pptx：{ext}")
    try:
        import win32com.client  # type: ignore
    except Exception as exc:  # noqa: BLE001
        raise ValueError(
            f"未安装 pywin32，无法直接解析 {ext}。请 `pip install pywin32` 并确保本机已安装 Microsoft Word/PowerPoint，"
            f"或先把 {ext} 手动另存为 {target_ext} 再上传。"
        ) from exc
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=target_ext)
    tmp.close()
    try:
        if target_ext == ".docx":
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            try:
                doc = word.Documents.Open(os.path.abspath(path))
                doc.SaveAs(tmp.name, FileFormat=16)  # 16 = wdFormatDocumentDefault (.docx)
                doc.Close(False)
            finally:
                word.Quit()
        else:
            ppt = win32com.client.DispatchEx("PowerPoint.Application")
            try:
                pres = ppt.Presentations.Open(os.path.abspath(path), WithWindow=False)
                pres.SaveAs(tmp.name, 24)  # 24 = ppSaveAsOpenXMLPresentation (.pptx)
                pres.Close()
            finally:
                try:
                    ppt.Quit()
                except Exception:  # noqa: BLE001
                    pass
    except Exception as exc:  # noqa: BLE001
        try:
            os.unlink(tmp.name)
        except OSError:
            pass
        raise ValueError(
            f"本机 Office 转换 {ext}→{target_ext} 失败：{exc}。"
            f"请确认 Microsoft Office 已正确安装并激活，或改用转换后的 {target_ext}。"
        ) from exc
    return tmp.name


def _read_ole_doc_fallback(path: str) -> str:
    """不依赖 Office：用 olefile 读 .doc 的 WordDocument 流文本（效果有限，仅兜底）。"""
    try:
        import olefile  # type: ignore
    except Exception as exc:  # noqa: BLE001
        raise ValueError("未安装 olefile，无法解析 .doc。请 `pip install olefile` 或安装 Microsoft Word。") from exc
    if not olefile.isOleFile(path):
        raise ValueError("非 OLE 文件，不是有效的 .doc。")
    with olefile.OleFileIO(path) as ole:
        if not ole.exists("WordDocument"):
            raise ValueError("不包含 WordDocument 流。")
        raw = ole.openstream("WordDocument").read()
    # 以 utf-16-le 解码纯文本段（会夹杂元数据，但规则标签足以应付）
    try:
        text = raw.decode("utf-16-le", errors="ignore")
    except Exception:  # noqa: BLE001
        text = raw.decode("gb18030", errors="ignore")
    text = re.sub(r"[\x00-\x08\x0b-\x1f\x7f]", "", text)
    text = re.sub(r"\s{3,}", "\n", text)
    return text.strip()


def _read_ole_ppt_fallback(path: str) -> str:
    """不依赖 Office：用 olefile 读 .ppt 的 User Strings / Current User Streams（效果有限）。"""
    try:
        import olefile  # type: ignore
    except Exception as exc:  # noqa: BLE001
        raise ValueError("未安装 olefile，无法解析 .ppt。请 `pip install olefile` 或安装 Microsoft PowerPoint。") from exc
    if not olefile.isOleFile(path):
        raise ValueError("非 OLE 文件，不是有效的 .ppt。")
    texts = []
    with olefile.OleFileIO(path) as ole:
        for st in ole.listdir():
            name = "/".join(st)
            low = name.lower()
            if any(k in low for k in ("current user", "user strings", "powerpoint document", "utf8", "unicode")):
                try:
                    raw = ole.openstream(st).read()
                except Exception:  # noqa: BLE001
                    continue
                for enc in ("utf-16-le", "utf-8", "gb18030"):
                    try:
                        t = raw.decode(enc, errors="ignore")
                    except Exception:  # noqa: BLE001
                        continue
                    t = re.sub(r"[\x00-\x08\x0b-\x1f\x7f]", " ", t)
                    t = re.sub(r" {2,}", " ", t).strip()
                    if len(t) > 30:
                        texts.append(t)
                        break
    joined = "\n".join(t for t in texts if t)
    if len(joined) < 40:
        raise ValueError("PPT OLE 兜底提取内容过少，请安装 Microsoft PowerPoint 以获得完整内容。")
    return joined


def parse_doc(path: str) -> ParsedDoc:
    """解析 .doc：优先 Office→docx，其次 olefile 兜底文本，都失败时降级为元数据占位条目（方案C）。"""
    # 路径 1：Office 转 docx → 复用 parse_docx
    tmp_path: str | None = None
    try:
        tmp_path = _office_to_temp(path, ".docx")
        doc = parse_docx(tmp_path)
        return doc
    except ValueError as first:
        logger.warning("Word 转换失败，尝试 olefile 兜底：%s", first)
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.unlink(tmp_path)
            except OSError:
                pass
    # 路径 2：olefile 兜底
    try:
        text = _read_ole_doc_fallback(path)
        return ParsedDoc(text=text, outline=[], pages={1: text}, ocr_used=False)
    except ValueError as second:
        logger.warning("olefile doc 兜底失败，降级为「元数据占位条目」入库：%s", second)
        return metadata_fallback_parsedoc(path, "doc")


def parse_ppt(path: str) -> ParsedDoc:
    """解析 .ppt：优先 Office→pptx，其次 olefile 兜底，都失败时降级为元数据占位条目（方案C）。"""
    tmp_path: str | None = None
    try:
        tmp_path = _office_to_temp(path, ".pptx")
        doc = parse_pptx(tmp_path)
        return doc
    except ValueError as first:
        logger.warning("PowerPoint 转换失败，尝试 olefile 兜底：%s", first)
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.unlink(tmp_path)
            except OSError:
                pass
    try:
        text = _read_ole_ppt_fallback(path)
        return ParsedDoc(text=text, outline=[], pages={1: text}, ocr_used=False)
    except ValueError as second:
        logger.warning("olefile ppt 兜底失败，降级为「元数据占位条目」入库：%s", second)
        return metadata_fallback_parsedoc(path, "ppt")


PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx, ".doc": parse_doc,
    ".pptx": parse_pptx, ".ppt": parse_ppt,
    ".txt": parse_plain, ".md": parse_plain, ".markdown": parse_plain,
    ".text": parse_plain,
    ".jpg": parse_image, ".jpeg": parse_image, ".png": parse_image,
}


def parse_file(path: str) -> ParsedDoc:
    ext = os.path.splitext(path)[1].lower()
    parser = PARSERS.get(ext)
    if parser is None:
        raise ValueError(f"暂不支持的文件格式：{ext}")
    doc = parser(path)
    # 统一最终兜底：任何解析结果 text 几乎为空时，都走元数据占位，保证上层不出现"课程=未分类"的最兜底形式
    if doc is None:
        doc = ParsedDoc()
    if len(re.sub(r"\s+", "", doc.text or "")) < 40 and ext in {".pdf", ".docx", ".pptx", ".txt", ".md"}:
        logger.warning("parse_file 最终文本极少，走统一元数据占位：%s", os.path.basename(path))
        kind = {"pdf": "pdf_scan", "docx": "docx", "pptx": "pptx"}.get(ext, ext.lstrip("."))
        doc = metadata_fallback_parsedoc(path, kind)
    return doc


# ---------------- 分块 ----------------

def chunk_document(parsed: ParsedDoc, max_len: int = 600, overlap: int = 80) -> list[dict]:
    """按章节切块：无章节结构时按滑动窗口。每块 {text, page, section}。"""
    chunks: list[dict] = []
    positions = []
    # 以全文偏移为锚，把 outline 变成区间
    search_from = 0
    for item in parsed.outline:
        idx = parsed.text.find(item["title"], search_from)
        if idx < 0:
            continue
        positions.append((idx, item["section"]))
        search_from = idx + 1
    positions.append((len(parsed.text), None))

    for i in range(len(positions) - 1):
        start, section = positions[i]
        end = positions[i + 1][0]
        seg = parsed.text[start:end].strip()
        if not seg:
            continue
        for piece in _split_long(seg, max_len, overlap):
            chunks.append({"text": piece, "section": section or "", "page": parsed.page_of(start)})
    if not chunks:  # 无任何章节结构
        for piece in _split_long(parsed.text, max_len, overlap):
            offset = parsed.text.find(piece[:30])
            chunks.append({"text": piece, "section": "", "page": parsed.page_of(max(offset, 0))})
    return chunks


def _split_long(text: str, max_len: int, overlap: int) -> list[str]:
    text = text.strip()
    if len(text) <= max_len:
        return [text]
    pieces, start = [], 0
    while start < len(text):
        end = min(start + max_len, len(text))
        # 尽量在句号/换行处断开
        if end < len(text):
            for punct in ("。", "\n", "；", "！", "？", ". "):
                cut = text.rfind(punct, start + max_len // 2, end)
                if cut > 0:
                    end = cut + 1
                    break
        pieces.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return [p for p in pieces if p]
